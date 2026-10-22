"""carrier_docs.py — dựng tài liệu chở kênh K1 ở ba định dạng MD/DOCX/PDF (STT-44).

LÝ DO TỒN TẠI
-------------
Injector K1 (`injector_k1.py`) cần một tài liệu doanh nghiệp vô hại có NHÚNG payload, để thả vào
kho tri thức RAG. Module này chỉ dựng tài liệu đó — nhận chuỗi, trả bytes/str — không biết gì về
Qdrant, retriever hay `delivered`. Tách khỏi injector để kiểm được việc dựng tài liệu mà không cần
dịch vụ vector store.

HAI LỐI NHÚNG
-------------
* **Hiển thị** (kỹ thuật T1/T2/T3/T5, và từng mảnh T6): chuỗi payload chèn thẳng vào thân tài liệu
  chở như một đoạn văn thường. Người duyệt thấy, parser đọc ra — không có bất đối xứng.
* **Ẩn** (T4): ủy thác cho `scripts/make_hidden_pdf.render_pdf`. Chỉ các bề mặt mà chuỗi ẩn rơi vào
  THÂN VĂN retriever đọc được mới dùng được ở K1 (xem `config/channel_render.yaml:K1`); việc ánh xạ
  obfuscation → bề mặt và chặn bề mặt không giao hàng được là của injector, không phải module này.

THÂN TÀI LIỆU CHỞ là nguồn chân lý duy nhất: dùng lại `CARRIER_TITLE`/`CARRIER_BODY` của
`make_hidden_pdf` để bản hiển thị và bản ẩn T4 cùng một bối cảnh, không đẻ hai thân khác nhau.

TÍNH TÁI LẬP: PDF sinh byte-identical (ngày tạo cố định, font subset theo thứ tự ký tự cố định),
giống `make_hidden_pdf`/`gen_corpus_pdfs`. DOCX của python-docx KHÔNG byte-identical (có timestamp
nội bộ); tái lập của K1 dựa trên NỘI DUNG trích ra ổn định, không trên byte của .docx.

KIỂM: tests/test_injector_k1.py.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import docx
from fpdf import FPDF

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(REPO_ROOT / "scripts"))

# Thân tài liệu chở và PDF helper dùng chung với bộ tạo tài liệu ẩn T4 — một nguồn, không chép lại.
from make_hidden_pdf import (  # noqa: E402
    CARRIER_BODY,
    CARRIER_TITLE,
    FIXED_CREATION_DATE,
    FONT_PATH,
)
from make_hidden_pdf import (  # noqa: E402
    render_pdf as render_hidden_pdf,
)

FORMATS = ("md", "docx", "pdf")
#: Đuôi file theo định dạng — khớp _READERS của rag.ingest.
SUFFIX = {"md": ".md", "docx": ".docx", "pdf": ".pdf"}


class CarrierDocError(ValueError):
    """Không dựng được tài liệu chở (định dạng lạ, vị trí nhúng lạ)."""


def _body_with_payload(payload: str, position: str) -> list[str]:
    """Thân tài liệu chở có chèn `payload` theo `position`. Trả danh sách dòng."""
    if position == "after_intro":
        # Sau dòng mở đầu (lời chào) — đoạn payload lọt cùng chunk với ngữ cảnh thật phía sau.
        return [CARRIER_BODY[0], payload, *CARRIER_BODY[1:]]
    raise CarrierDocError(f"visible_position lạ '{position}'; hiện chỉ hỗ trợ 'after_intro'")


def build_visible_md(payload: str, *, position: str = "after_intro") -> bytes:
    """Markdown chở payload hiển thị. Trả bytes UTF-8 (ghi thẳng ra file)."""
    lines = [f"# {CARRIER_TITLE}", "", *_body_with_payload(payload, position)]
    return ("\n\n".join(lines) + "\n").encode("utf-8")


def build_visible_docx(payload: str, *, position: str = "after_intro") -> bytes:
    """DOCX chở payload hiển thị. Trả bytes của file .docx."""
    document = docx.Document()
    document.add_heading(CARRIER_TITLE, level=1)
    for line in _body_with_payload(payload, position):
        document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_visible_pdf(payload: str, *, position: str = "after_intro") -> bytes:
    """PDF chở payload hiển thị (chữ đen cỡ thường). Tái lập byte-identical.

    Dùng cùng font/ngày cố định như make_hidden_pdf để byte ổn định giữa các lần chạy.
    """
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_creation_date(FIXED_CREATION_DATE)
    pdf.set_title(CARRIER_TITLE)
    pdf.add_font("DejaVu", "", str(FONT_PATH))
    pdf.set_margins(left=20, top=20, right=20)
    pdf.add_page()
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("DejaVu", size=15)
    pdf.multi_cell(0, 9, CARRIER_TITLE, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("DejaVu", size=11)
    for line in _body_with_payload(payload, position):
        pdf.multi_cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


_VISIBLE_BUILDERS = {
    "md": build_visible_md,
    "docx": build_visible_docx,
    "pdf": build_visible_pdf,
}


def build_visible(payload: str, fmt: str, *, position: str = "after_intro") -> bytes:
    """Tài liệu chở payload HIỂN THỊ ở định dạng `fmt` ∈ FORMATS. Trả bytes."""
    builder = _VISIBLE_BUILDERS.get(fmt)
    if builder is None:
        raise CarrierDocError(f"định dạng lạ '{fmt}'; cho {FORMATS}")
    return builder(payload, position=position)


def build_hidden_t4(instruction: str, surface: str) -> bytes:
    """Tài liệu chở payload T4 ẩn theo `surface` (PDF). Trả bytes.

    `instruction` là chỉ thị TRẦN (không phải base_text đã bọc technique.template): bề mặt T4 tự
    là lớp obfuscation, bọc thêm template của payload sẽ ẩn hai lần. Chỉ nhận bề mặt PDF — html
    và metadata không giao được qua retriever (xem docstring module và config K1).
    """
    return render_hidden_pdf(instruction, surface)


def write_doc(out_dir: Path | str, doc_id: str, fmt: str, data: bytes) -> Path:
    """Ghi `data` ra `<out_dir>/<doc_id><suffix>`. Trả đường dẫn đã ghi."""
    suffix = SUFFIX.get(fmt)
    if suffix is None:
        raise CarrierDocError(f"định dạng lạ '{fmt}'; cho {FORMATS}")
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    # Thay ký tự lạ trong doc_id (có '::') để thành tên file hợp lệ.
    path = out / f"{doc_id.replace('::', '__').replace('/', '_')}{suffix}"
    path.write_bytes(data)
    return path
